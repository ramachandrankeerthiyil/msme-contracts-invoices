"""Reading an uploaded invoice workbook (INV-001 design "Logic §1–§3").

File-level problems raise AppError, which means nothing is saved (AC2, AC3).
"""

import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.core.errors import AppError
from app.domain.invoice_rules import OPTIONAL_COLUMNS, REQUIRED_COLUMNS

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20_000

_ZIP_MAGIC = b"PK\x03\x04"
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # legacy .xls, or a password-protected .xlsx
_ENCRYPTION_STREAM = "EncryptionInfo".encode("utf-16-le")


def invalid_file_type() -> AppError:
    return AppError("INVALID_FILE_TYPE", "Please upload an Excel file (.xlsx).")


def file_too_large() -> AppError:
    return AppError(
        "FILE_TOO_LARGE",
        "This file is larger than 10 MB. Please upload a smaller file.",
        status=413,
        details={"max_bytes": MAX_UPLOAD_BYTES},
    )


def check_file_type(path: Path) -> None:
    """Checks the content, not the extension: a renamed PDF is still a PDF."""
    with path.open("rb") as fh:
        head = fh.read(8)
    if head.startswith(_ZIP_MAGIC):
        try:
            with zipfile.ZipFile(path) as archive:
                if "xl/workbook.xml" in archive.namelist():
                    return
        except zipfile.BadZipFile:
            pass
        raise invalid_file_type()
    if head.startswith(_OLE_MAGIC) and _ENCRYPTION_STREAM in path.read_bytes():
        raise AppError(
            "FILE_PROTECTED",
            "This file is password-protected. Please remove the password and try again.",
        )
    raise invalid_file_type()


def normalise_header(value: Any) -> str:
    return " ".join(str(value).split()).lower() if value is not None else ""


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


@dataclass(frozen=True)
class RawRow:
    row_number: int  # the Excel row number the user sees
    values: dict[str, Any]
    formula_columns: frozenset[str] = field(default_factory=frozenset)


def read_rows(path: Path) -> list[RawRow]:
    """Reads the first worksheet. Blank rows are skipped and not counted."""
    try:
        # Loaded twice: cached values for the data, formulas to spot formulas with no saved value.
        values_book = load_workbook(path, read_only=True, data_only=True)
        formulas_book = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:  # corrupt or unsupported workbook
        raise invalid_file_type() from exc

    try:
        value_rows = values_book.worksheets[0].iter_rows(values_only=True)
        formula_rows = formulas_book.worksheets[0].iter_rows(values_only=True)
        header = next(value_rows, None) or ()
        next(formula_rows, None)

        columns: dict[str, int] = {}
        for index, name in enumerate(header):
            columns.setdefault(normalise_header(name), index)
        missing = [c for c in REQUIRED_COLUMNS if c.lower() not in columns]
        if missing:
            raise AppError(
                "MISSING_COLUMNS",
                f"Your file is missing these columns: {', '.join(missing)}. "
                "Download the template to see the correct layout.",
                status=422,
                details={"missing": missing, "found": [str(h) for h in header if h is not None]},
            )
        index_of = {column: columns[column.lower()] for column in REQUIRED_COLUMNS}
        index_of.update(
            {
                column: columns[column.lower()]
                for column in OPTIONAL_COLUMNS
                if column.lower() in columns
            }
        )

        rows: list[RawRow] = []
        for row_number, (cells, formulas) in enumerate(
            zip(value_rows, formula_rows, strict=True), start=2
        ):
            values = {c: cells[i] if i < len(cells) else None for c, i in index_of.items()}
            raw = {c: formulas[i] if i < len(formulas) else None for c, i in index_of.items()}
            formula_columns = frozenset(
                c
                for c in REQUIRED_COLUMNS
                if values[c] is None and isinstance(raw[c], str) and raw[c].startswith("=")
            )
            if not formula_columns and all(_is_blank(v) for v in values.values()):
                continue
            rows.append(RawRow(row_number, values, formula_columns))
            if len(rows) > MAX_ROWS:
                raise AppError(
                    "TOO_MANY_ROWS",
                    f"This file has more than {MAX_ROWS:,} rows. "
                    "Please split it into smaller files.",
                    status=422,
                    details={"max_rows": MAX_ROWS},
                )
    finally:
        values_book.close()
        formulas_book.close()

    if not rows:
        raise AppError(
            "NO_DATA_ROWS", "We couldn't find any invoice rows below the header.", status=422
        )
    return rows
