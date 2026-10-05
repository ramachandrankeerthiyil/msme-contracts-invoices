"""The downloadable invoice template (INV-001 AC11). Only the first sheet is read on upload."""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.domain.invoice_rules import OPTIONAL_COLUMNS, REQUIRED_COLUMNS

_HEADER_FILL = PatternFill("solid", fgColor="D6E8F9")
_HEADER_FONT = Font(bold=True, color="0F2742")

INSTRUCTIONS = [
    (
        "Invoice Number",
        "Required. Your invoice reference, e.g. INV-1001. Uploading the same number "
        "again updates that invoice.",
    ),
    ("Customer Name", "Required. Who the invoice was sent to."),
    ("Date Raised", "Required. The invoice date, e.g. 05-Sep-2026 or 05/09/2026 (day first)."),
    ("Due Date", "Required. When payment is due. Must be on or after the Date Raised."),
    (
        "Amount",
        "Required. The invoice total in rupees, e.g. 125000 or 1,25,000.50. "
        "Must be more than zero, with at most 2 decimal places.",
    ),
    (
        "Paid Date",
        "Leave blank if the invoice is not paid yet. Fill in the payment date once paid.",
    ),
    (
        "Customer Email",
        "Optional. The email address of the person or team to remind about this invoice, "
        "e.g. accounts@yourclient.in. Needed to send an email reminder. "
        "You can leave this column out completely.",
    ),
]


def build_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Invoices"
    sheet.append([*REQUIRED_COLUMNS, *OPTIONAL_COLUMNS])
    sheet.append(
        [
            "INV-1001",
            "Example Traders Pvt Ltd",
            date(2026, 9, 1),
            date(2026, 10, 1),
            125000.00,
            None,
            "accounts@example.com",
        ]
    )
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    for column in "CDF":
        sheet[f"{column}2"].number_format = "DD-MMM-YYYY"
    sheet["E2"].number_format = "#,##,##0.00"
    for column, width in zip("ABCDEFG", [18, 32, 15, 15, 15, 15, 32], strict=True):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"

    help_sheet = workbook.create_sheet("How to fill this in")
    help_sheet.append(["Column", "What to enter"])
    for row in INSTRUCTIONS:
        help_sheet.append(list(row))
    help_sheet.append([])
    help_sheet.append(
        ["Note", "Put your invoices on the first sheet. Only the first sheet is read."]
    )
    for cell in help_sheet[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    help_sheet.column_dimensions["A"].width = 18
    help_sheet.column_dimensions["B"].width = 90
    for row in help_sheet.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
