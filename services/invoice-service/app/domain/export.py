"""Export the current invoice view to Excel (INV-002 AC9)."""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.db.invoice_queries import InvoiceRow
from app.domain.clock import local_date
from app.domain.invoice_status import STATUS_LABELS, due_hint

MAX_EXPORT_ROWS = 20_000
HEADERS = [
    "Invoice Number",
    "Customer Name",
    "Date Raised",
    "Due Date",
    "Amount",
    "Paid Date",
    "Status",
    "Due",
    "Record status",
]


def record_status_label(row: InvoiceRow, time_zone: str) -> str:
    updated_at = row.invoice.record_updated_at
    if row.invoice.record_status == "updated" and updated_at is not None:
        updated = local_date(updated_at, time_zone)
        return f"Updated on {updated.day} {updated:%b %Y}"
    return "New"


def build_export(rows: list[InvoiceRow], *, today: date, time_zone: str) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Invoices"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="0F2742")
        cell.fill = PatternFill("solid", fgColor="D6E8F9")

    for row in rows:
        invoice = row.invoice
        sheet.append(
            [
                invoice.invoice_number,
                invoice.customer_name,
                invoice.date_raised,
                invoice.due_date,
                invoice.amount,
                invoice.paid_date,
                STATUS_LABELS[row.status],
                due_hint(invoice.due_date, invoice.paid_date, today),
                record_status_label(row, time_zone),
            ]
        )
        current = sheet.max_row
        for column in "CDF":
            sheet[f"{column}{current}"].number_format = "DD-MMM-YYYY"
        sheet[f"E{current}"].number_format = "#,##,##0.00"

    for column, width in zip("ABCDEFGHI", [16, 32, 14, 14, 15, 14, 14, 20, 24], strict=True):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
