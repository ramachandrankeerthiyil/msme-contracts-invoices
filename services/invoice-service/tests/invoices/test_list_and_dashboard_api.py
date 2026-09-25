"""INV-002 (list, export) and INV-003 (dashboard) API tests, with "today" pinned to 25 Sep 2026."""

import asyncio
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from urllib.parse import urlencode

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.domain.clock import get_today
from app.domain.invoice_status import status_of
from app.main import create_app
from tests.conftest import _execute
from tests.invoices.workbooks import row, upload, workbook_bytes

TODAY = date(2026, 9, 25)
UPLOAD_DAY = date(2026, 9, 24)  # U: the dashboard week is 24–30 Sep


def d(offset: int) -> date:
    return TODAY + timedelta(days=offset)


# Mirrors samples/invoices-sample.xlsx as seen on 25 Sep (generated 24 Sep).
SAMPLE = [
    row("INV-2601", "Sharma Textiles Pvt Ltd", d(-76), d(-46), 185000, d(-51)),
    row("INV-2602", "Kaveri Agro Foods", d(-61), d(-31), 42750.50, d(-29)),
    row("INV-2603", "Bluewave Logistics LLP", d(-41), d(-11), 96400, d(-13)),
    row("INV-2604", "Mehta Auto Components", d(-21), d(9), 310000, d(-3)),
    row("INV-2605", "Green Leaf Organics", d(-16), d(1), 18500, d(-2)),
    row("INV-2606", "Deccan Printing Works", d(-71), d(-41), 125000, None),
    row("INV-2607", "Sunrise Pharma Distributors", d(-46), d(-16), 267500.75, None),
    row("INV-2608", "Nilgiri Tea Traders", d(-36), d(-6), 54000, None),
    row("INV-2609", "Kaveri Agro Foods", d(-32), d(-2), 8900, None),
    row("INV-2610", "Mehta Auto Components", d(-31), d(-1), 142000, None),
    row("INV-2611", "Bluewave Logistics LLP", d(-29), d(0), 76250, None),
    row("INV-2612", "Sharma Textiles Pvt Ltd", d(-28), d(2), 225000, None),
    row("INV-2613", "Green Leaf Organics", d(-26), d(4), 31500, None),
    row("INV-2614", "Deccan Printing Works", d(-21), d(5), 64000, None),
    row("INV-2615", "Sunrise Pharma Distributors", d(-11), d(19), 410000, None),
    row("INV-2616", "Nilgiri Tea Traders", d(-6), d(24), 22800, None),
    row("INV-2617", "Kaveri Agro Foods", d(-3), d(44), 98000, None),
]


@pytest.fixture
def client(settings, uploads_dir, clean_tables):
    app = create_app(settings.model_copy(update={"uploads_dir": str(uploads_dir)}))
    app.dependency_overrides[get_today] = lambda: TODAY
    with TestClient(app) as test_client:
        yield test_client


def set_upload_day(settings, day: date) -> None:
    """Uploads are stamped with the real clock; pin them to a known local date."""
    asyncio.run(
        _execute(
            settings,
            f'UPDATE "{settings.db_schema}".invoice_uploads '
            f"SET uploaded_at = TIMESTAMPTZ '{day.isoformat()} 10:00:00+05:30'",
        )
    )


@pytest.fixture
def sample(client, settings):
    assert upload(client, workbook_bytes(SAMPLE)).json()["rows_created"] == 17
    set_upload_day(settings, UPLOAD_DAY)
    return client


def get_list(client, **params):
    response = client.get(f"/api/invoices?{urlencode(params)}")
    assert response.status_code == 200, response.text
    return response.json()


def numbers(body):
    return [item["invoice_number"] for item in body["items"]]


# --- INV-002 --------------------------------------------------------------------------------


def test_INV_002_AC2_every_status_matches_the_rule(sample):
    body = get_list(sample, view="all", page_size=100)

    for item in body["items"]:
        due = date.fromisoformat(item["due_date"])
        paid = date.fromisoformat(item["paid_date"]) if item["paid_date"] else None
        assert item["status"] == status_of(due, paid, TODAY, 5).value, item["invoice_number"]
    assert body["today"] == "2026-09-25"


def test_INV_002_AC3_counts_per_view(sample):
    body = get_list(sample)

    assert body["counts"] == {
        "follow_up": 9,
        "outstanding": 5,
        "at_risk": 4,
        "open": 3,
        "paid": 5,
        "all": 17,
    }
    assert body["total"] == 9  # default view is follow_up


def test_INV_002_AC3_counts_ignore_the_view_but_respect_search(sample):
    assert get_list(sample, view="paid")["counts"]["all"] == 17
    assert get_list(sample, q="kaveri")["counts"] == {
        "follow_up": 1,
        "outstanding": 1,
        "at_risk": 0,
        "open": 1,
        "paid": 1,
        "all": 3,
    }


def test_INV_002_AC4_default_order_outstanding_most_overdue_then_at_risk_soonest(sample):
    body = get_list(sample)

    assert numbers(body) == [
        # outstanding, most overdue first
        *["INV-2606", "INV-2607", "INV-2608", "INV-2609", "INV-2610"],
        # at risk, soonest due first
        *["INV-2611", "INV-2612", "INV-2613", "INV-2614"],
    ]


@pytest.mark.parametrize(
    ("sort", "first_asc"),
    [
        ("invoice_number", "INV-2601"),
        ("customer_name", "INV-2603"),  # Bluewave …
        ("amount", "INV-2609"),  # 8,900
        ("due_date", "INV-2601"),
        ("date_raised", "INV-2601"),
        ("paid_date", "INV-2601"),
        ("status", "INV-2606"),
    ],
)
def test_INV_002_AC4_every_column_sorts(sample, sort, first_asc):
    ascending = numbers(get_list(sample, view="all", sort=sort, order="asc", page_size=100))

    assert ascending[0] == first_asc
    assert len(ascending) == len(set(ascending)) == 17


@pytest.mark.parametrize("sort", ["invoice_number", "amount", "due_date"])  # unique values
def test_INV_002_AC4_descending_is_the_reverse(sample, sort):
    ascending = numbers(get_list(sample, view="all", sort=sort, order="asc", page_size=100))
    descending = numbers(get_list(sample, view="all", sort=sort, order="desc", page_size=100))

    assert descending == list(reversed(ascending))


def test_INV_002_AC4_unpaid_invoices_sort_after_paid_ones_by_paid_date(sample):
    for order in ("asc", "desc"):
        items = get_list(sample, view="all", sort="paid_date", order=order, page_size=100)["items"]
        assert all(item["paid_date"] is None for item in items[5:])  # nulls last both ways


def test_INV_002_AC5_search_by_number_or_customer_case_insensitive(sample):
    assert numbers(get_list(sample, view="all", q="inv-2607")) == ["INV-2607"]
    assert set(numbers(get_list(sample, view="all", q="GREEN leaf"))) == {"INV-2605", "INV-2613"}
    assert get_list(sample, view="all", q="100%_sure")["total"] == 0  # LIKE wildcards are literal


def test_INV_002_AC5_updated_only(sample, settings):
    changed = row("INV-2606", "Deccan Printing Works", d(-71), d(-41), 130000)
    upload(sample, workbook_bytes([changed]))

    body = get_list(sample, view="all", updated="true")

    assert numbers(body) == ["INV-2606"]
    assert body["items"][0]["record_status"] == "updated"
    assert body["items"][0]["record_updated_at"] is not None


def test_INV_002_AC6_total_amount_covers_all_pages(sample):
    body = get_list(sample, page_size=2)

    assert len(body["items"]) == 2
    assert body["total"] == 9
    assert Decimal(body["total_amount"]) == Decimal("994150.75")  # 5 outstanding + 4 at risk
    assert body["items"][0]["amount"] == "125000.00"


def test_INV_002_paging_and_validation(sample):
    page2 = get_list(sample, view="all", sort="invoice_number", page=2, page_size=5)
    assert numbers(page2) == ["INV-2606", "INV-2607", "INV-2608", "INV-2609", "INV-2610"]
    assert get_list(sample, view="all", page=99)["items"] == []

    assert sample.get("/api/invoices?view=nope").status_code == 422
    assert sample.get("/api/invoices?page_size=500").status_code == 422
    bad_range = sample.get("/api/invoices?due_from=2026-10-01&due_to=2026-09-01")
    assert bad_range.status_code == 422
    message = bad_range.json()["error"]["message"]
    assert message == "The start date must be on or before the end date."


def test_INV_002_AC9_export_matches_the_current_view(sample):
    response = sample.get("/api/invoices/export?view=outstanding")

    assert response.status_code == 200
    assert 'filename="invoices-2026-09-25.xlsx"' in response.headers["content-disposition"]
    sheet = load_workbook(BytesIO(response.content)).worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == (
        "Invoice Number",
        "Customer Name",
        "Date Raised",
        "Due Date",
        "Amount",
        "Paid Date",
        "Status",
        "Due",
        "Record status",
    )
    assert [r[0] for r in rows[1:]] == ["INV-2606", "INV-2607", "INV-2608", "INV-2609", "INV-2610"]
    assert rows[1][6:] == ("Outstanding", "41 days overdue", "New")


def test_PLT_002_AC7_prefixed_routes_get_full_templates_in_metrics(sample):
    from app.core.metrics import REGISTRY

    def count(route):
        labels = {"method": "GET", "route": route, "status": "200"}
        return REGISTRY.get_sample_value("http_requests_total", labels) or 0.0

    before_list, before_dash = count("/api/invoices"), count("/api/invoices/dashboard")
    sample.get("/api/invoices")
    sample.get("/api/invoices/dashboard")

    assert count("/api/invoices") == before_list + 1
    assert count("/api/invoices/dashboard") == before_dash + 1


def test_INV_002_empty_database(client):
    body = get_list(client)

    assert body["items"] == []
    assert body["counts"]["all"] == 0
    assert body["total_amount"] == "0.00"


# --- INV-003 --------------------------------------------------------------------------------


def dashboard(client):
    response = client.get("/api/invoices/dashboard")
    assert response.status_code == 200, response.text
    return response.json()


def test_INV_003_AC9_no_uploads_means_no_data(client):
    assert dashboard(client) == {"today": "2026-09-25", "has_data": False}


def test_INV_003_AC1_week_is_seven_days_from_the_latest_completed_upload(sample, settings):
    body = dashboard(sample)
    assert body["week"]["start"] == "2026-09-24"
    assert body["week"]["end"] == "2026-09-30"

    # A later upload where every row was rejected does not move the week.
    upload(sample, workbook_bytes([row("BAD", amount=0)]))
    assert dashboard(sample)["week"]["start"] == "2026-09-24"


def test_INV_003_AC2_AC3_value_and_count_this_week(sample):
    body = dashboard(sample)

    # Due 24–30 Sep: INV-2605 (paid), 2610 (outstanding), 2611–2614 (at risk)
    assert body["invoices_this_week"] == 6
    assert body["value_this_week"] == "557250.00"


def test_INV_003_AC4_follow_up_includes_invoices_overdue_before_the_week(sample):
    body = dashboard(sample)

    assert body["follow_up"] == {"total": 9, "outstanding": 5, "at_risk": 4}


def test_INV_003_AC6_every_card_link_returns_exactly_its_number(sample):
    body = dashboard(sample)

    this_week = get_list(sample, **body["links"]["this_week"])
    assert this_week["total"] == body["invoices_this_week"]
    assert this_week["total_amount"] == body["value_this_week"]
    follow_up = get_list(sample, **body["links"]["follow_up"])
    assert follow_up["total"] == body["follow_up"]["total"]


def test_INV_003_AC7_top_five_most_overdue_first(sample):
    top = dashboard(sample)["top_follow_up"]

    assert [item["invoice_number"] for item in top] == [
        "INV-2606",
        "INV-2607",
        "INV-2608",
        "INV-2609",
        "INV-2610",
    ]
    assert top[0]["days_until_due"] == -41
    assert top[0]["amount"] == "125000.00"


def test_INV_003_AC8_value_by_status_always_lists_all_four(sample):
    assert dashboard(sample)["value_by_status"] == [
        {"status": "outstanding", "amount": "142000.00", "count": 1},
        {"status": "at_risk", "amount": "396750.00", "count": 4},
        {"status": "open", "amount": "0.00", "count": 0},
        {"status": "paid", "amount": "18500.00", "count": 1},
    ]
