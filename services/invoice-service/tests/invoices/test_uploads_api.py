import time
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.core.metrics import REGISTRY
from app.domain import upload_service
from tests.invoices.workbooks import HEADERS, HEADERS_WITH_EMAIL, row, upload, workbook_bytes

# --- AC2: file checks — nothing saved ------------------------------------------------------


def _assert_nothing_saved(db_rows, uploads_dir):
    assert db_rows("SELECT id FROM invoice_uploads") == []
    assert db_rows("SELECT id FROM invoices") == []
    assert not uploads_dir.exists() or not any(uploads_dir.iterdir())


@pytest.mark.parametrize(
    ("content", "code"),
    [
        (b"%PDF-1.4 not really a spreadsheet", "INVALID_FILE_TYPE"),
        (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512, "INVALID_FILE_TYPE"),  # legacy .xls
        (
            b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
            + b"\x00" * 64
            + "EncryptionInfo".encode("utf-16-le"),
            "FILE_PROTECTED",
        ),
        (b"PK\x03\x04 a zip that is not a workbook", "INVALID_FILE_TYPE"),
        (b"Invoice Number,Customer Name\nINV-1,Acme\n", "INVALID_FILE_TYPE"),  # CSV renamed
    ],
)
def test_INV_001_AC2_wrong_file_types_are_rejected(
    upload_client, db_rows, uploads_dir, content, code
):
    response = upload(upload_client, content)

    assert response.status_code == 400
    assert response.json()["error"]["code"] == code
    _assert_nothing_saved(db_rows, uploads_dir)


def test_INV_001_AC2_files_over_10_mb_are_rejected(upload_client, db_rows, uploads_dir):
    response = upload(upload_client, b"PK\x03\x04" + b"\x00" * (10 * 1024 * 1024))

    assert response.status_code == 413
    error = response.json()["error"]
    assert error["code"] == "FILE_TOO_LARGE"
    assert error["message"] == "This file is larger than 10 MB. Please upload a smaller file."
    _assert_nothing_saved(db_rows, uploads_dir)


# --- AC3: headers --------------------------------------------------------------------------


@pytest.mark.parametrize("missing", HEADERS)
def test_INV_001_AC3_each_missing_column_rejects_the_file(
    upload_client, db_rows, uploads_dir, missing
):
    headers = [h for h in HEADERS if h != missing]
    content = workbook_bytes([row()[: len(headers)]], headers=headers)

    response = upload(upload_client, content)

    assert response.status_code == 422
    error = response.json()["error"]
    assert error["code"] == "MISSING_COLUMNS"
    assert error["details"]["missing"] == [missing]
    assert missing in error["message"]
    _assert_nothing_saved(db_rows, uploads_dir)


def test_INV_001_AC3_header_case_spacing_order_and_extra_columns(upload_client):
    headers = [
        "  paid date",
        "AMOUNT",
        "Notes",
        "due   date",
        "date raised",
        "customer name",
        "Invoice Number",
    ]
    content = workbook_bytes(
        [[None, 500, "ignored", date(2026, 10, 1), date(2026, 9, 1), "Acme", "INV-9"]],
        headers=headers,
    )

    response = upload(upload_client, content)

    assert response.status_code == 201
    assert response.json()["rows_created"] == 1


def test_INV_001_header_only_file_has_no_data_rows(upload_client):
    response = upload(upload_client, workbook_bytes([]))

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "NO_DATA_ROWS"


# --- AC4, AC5, AC7, AC8: rows --------------------------------------------------------------


def test_INV_001_AC4_valid_rows_are_saved_alongside_rejected_ones(upload_client, db_rows):
    content = workbook_bytes(
        [
            row("INV-1"),
            row("INV-2", customer=None),
            [None, None, None, None, None, None],  # blank row: skipped, not counted
            row("INV-3", amount=-5),
            row("INV-4", paid=date(2026, 9, 20)),
        ]
    )

    response = upload(upload_client, content)

    assert response.status_code == 201
    body = response.json()
    assert body["status"] == "completed"
    assert (body["rows_total"], body["rows_created"], body["rows_rejected"]) == (4, 2, 2)
    assert body["rejections"] == [
        {"row": 3, "invoice_number": "INV-2", "reason": "Customer Name is missing."},
        {"row": 5, "invoice_number": "INV-3", "reason": "Amount must be greater than zero."},
    ]
    saved = db_rows("SELECT invoice_number, record_status, paid_date FROM invoices ORDER BY 1")
    assert saved == [
        {"invoice_number": "INV-1", "record_status": "new", "paid_date": None},
        {"invoice_number": "INV-4", "record_status": "new", "paid_date": date(2026, 9, 20)},
    ]


def test_INV_001_AC5_AC6_second_upload_overwrites_and_marks_updated(upload_client, db_rows):
    first = upload(upload_client, workbook_bytes([row("INV-1", amount=100), row("INV-2")])).json()
    second = upload(
        upload_client,
        workbook_bytes(
            [
                row("inv-1", customer="Acme Renamed", amount=250.5, paid=date(2026, 9, 30)),
                row("INV-3"),
            ]
        ),
    ).json()

    assert (first["rows_created"], first["rows_updated"]) == (2, 0)
    assert (second["rows_created"], second["rows_updated"]) == (1, 1)
    invoices = {r["invoice_number_key"]: r for r in db_rows("SELECT * FROM invoices")}
    updated = invoices["INV-1"]
    assert updated["invoice_number"] == "inv-1"
    assert updated["customer_name"] == "Acme Renamed"
    assert updated["amount"] == Decimal("250.50")
    assert updated["paid_date"] == date(2026, 9, 30)
    assert updated["record_status"] == "updated"
    assert updated["record_updated_at"] is not None
    assert str(updated["first_upload_id"]) == first["id"]
    assert str(updated["last_upload_id"]) == second["id"]
    assert invoices["INV-2"]["record_status"] == "new"
    assert invoices["INV-2"]["record_updated_at"] is None
    assert invoices["INV-3"]["record_status"] == "new"


def test_INV_001_AC6a_in_file_duplicate_last_row_wins(upload_client, db_rows):
    content = workbook_bytes([row("INV-1", amount=100), row("INV-2"), row("INV-1", amount=300)])

    body = upload(upload_client, content).json()

    assert body["overwrites"] == [{"row": 4, "replaced_row": 2, "invoice_number": "INV-1"}]
    assert (body["rows_total"], body["rows_created"], body["rows_overwritten"]) == (3, 2, 1)
    saved = db_rows(
        "SELECT amount, record_status, record_updated_at FROM invoices "
        "WHERE invoice_number_key = 'INV-1'"
    )[0]
    assert saved["amount"] == Decimal("300.00")
    assert saved["record_status"] == "updated"
    assert saved["record_updated_at"] is not None


def test_INV_001_AC6a_invalid_later_duplicate_does_not_overwrite(upload_client, db_rows):
    content = workbook_bytes([row("INV-1", amount=100), row("INV-1", amount=-1)])

    body = upload(upload_client, content).json()

    assert body["overwrites"] == []
    assert body["rows_rejected"] == 1
    saved = db_rows("SELECT amount, record_status FROM invoices")
    assert saved == [{"amount": Decimal("100.00"), "record_status": "new"}]


def test_INV_001_row_counts_always_add_up(upload_client):
    content = workbook_bytes([row("A"), row("B", amount=0), row("A"), row("C"), row("C"), row("D")])

    body = upload(upload_client, content).json()

    assert body["rows_total"] == (
        body["rows_created"]
        + body["rows_updated"]
        + body["rows_overwritten"]
        + body["rows_rejected"]
    )


def test_INV_001_AC8_every_upload_is_recorded_including_all_rejected(
    upload_client, db_rows, uploads_dir
):
    body = upload(upload_client, workbook_bytes([row(amount=0)]), name="bad.xlsx").json()

    assert body["status"] == "no_valid_rows"
    assert body["rows_rejected"] == 1
    recorded = db_rows("SELECT file_name, status, stored_path FROM invoice_uploads")
    assert recorded[0]["file_name"] == "bad.xlsx"
    assert recorded[0]["status"] == "no_valid_rows"
    assert (uploads_dir / f"{body['id']}.xlsx").exists()


def test_INV_001_AC8_upload_history_newest_first(upload_client):
    upload(upload_client, workbook_bytes([row("INV-1")]), name="first.xlsx")
    upload(upload_client, workbook_bytes([row("INV-2")]), name="second.xlsx")

    page = upload_client.get("/api/invoices/uploads").json()

    assert page["total"] == 2
    assert [item["file_name"] for item in page["items"]] == ["second.xlsx", "first.xlsx"]
    assert "rejections" not in page["items"][0]


def test_INV_001_formula_without_saved_value_is_rejected_with_explanation(upload_client):
    book = Workbook()
    sheet = book.active
    sheet.append(HEADERS)
    sheet.append(["INV-1", "Acme", date(2026, 9, 1), date(2026, 10, 1), "=100*2", None])
    buffer = BytesIO()
    book.save(buffer)  # openpyxl never stores computed values, like a file saved by a script

    body = upload(upload_client, buffer.getvalue()).json()

    assert body["rows_rejected"] == 1
    assert "Amount contains a formula with no saved value" in body["rejections"][0]["reason"]


# --- AC11: template ------------------------------------------------------------------------


def test_INV_001_AC11_template_has_the_exact_headers_and_uploads_cleanly(upload_client, db_rows):
    response = upload_client.get("/api/invoices/template")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert 'filename="invoice-template.xlsx"' in response.headers["content-disposition"]
    sheet = load_workbook(BytesIO(response.content)).worksheets[0]
    assert [cell.value for cell in sheet[1]] == HEADERS_WITH_EMAIL

    body = upload(upload_client, response.content).json()
    assert (body["rows_created"], body["rows_rejected"]) == (1, 0)
    assert db_rows("SELECT customer_email FROM invoices") == [
        {"customer_email": "accounts@example.com"}
    ]


# --- AC13: optional Customer Email ----------------------------------------------------------


def _emails(db_rows):
    rows = db_rows("SELECT invoice_number, customer_email FROM invoices ORDER BY invoice_number")
    return {r["invoice_number"]: r["customer_email"] for r in rows}


def _upload_with_email(client, rows):
    return upload(client, workbook_bytes(rows, headers=HEADERS_WITH_EMAIL)).json()


def test_INV_001_AC13_email_is_stored_trimmed_and_lower_cased(upload_client, db_rows):
    body = _upload_with_email(
        upload_client,
        [row("INV-1") + ["  Accounts@Acme.IN "], row("INV-2") + [None]],
    )

    assert (body["rows_created"], body["rows_rejected"]) == (2, 0)
    assert _emails(db_rows) == {"INV-1": "accounts@acme.in", "INV-2": None}


def test_INV_001_AC13_a_sheet_without_the_column_still_uploads(upload_client, db_rows):
    body = upload(upload_client, workbook_bytes([row("INV-1")])).json()

    assert (body["rows_created"], body["rows_rejected"]) == (1, 0)
    assert _emails(db_rows) == {"INV-1": None}


def test_INV_001_AC13_a_later_upload_overwrites_the_email_and_a_blank_clears_it(
    upload_client, db_rows
):
    _upload_with_email(upload_client, [row("INV-1") + ["old@acme.in"], row("INV-2") + ["x@y.in"]])

    body = _upload_with_email(
        upload_client, [row("INV-1") + ["new@acme.in"], row("INV-2") + [None]]
    )

    assert body["rows_updated"] == 2
    assert _emails(db_rows) == {"INV-1": "new@acme.in", "INV-2": None}


def test_INV_001_AC13_an_invalid_address_rejects_only_that_row(upload_client, db_rows):
    body = _upload_with_email(
        upload_client,
        [row("INV-1") + ["accounts@"], row("INV-2") + ["fine@acme.in"]],
    )

    assert (body["rows_created"], body["rows_rejected"]) == (1, 1)
    assert body["rejections"] == [
        {
            "row": 2,
            "invoice_number": "INV-1",
            "reason": "Customer Email 'accounts@' is not a valid email address.",
        }
    ]
    assert _emails(db_rows) == {"INV-2": "fine@acme.in"}


# --- AC12: all or nothing ------------------------------------------------------------------


def test_INV_001_AC12_failure_after_upsert_rolls_everything_back(
    upload_client, db_rows, uploads_dir, monkeypatch
):
    def fail_copy(*_args, **_kwargs):
        raise OSError("disk full")

    monkeypatch.setattr(upload_service.shutil, "copyfile", fail_copy)

    response = upload(upload_client, workbook_bytes([row("INV-1"), row("INV-2")]))

    assert response.status_code == 500
    assert response.json()["error"]["code"] == "INTERNAL_ERROR"
    _assert_nothing_saved(db_rows, uploads_dir)


# --- Observability (PLT-002 AC4, AC7) -------------------------------------------------------


def test_INV_001_logs_and_metrics(upload_client, logs):
    def sample(name, labels):
        return REGISTRY.get_sample_value(name, labels) or 0.0

    before_uploads = sample("invoice_uploads_total", {"result": "success"})
    before_rejected = sample("invoice_rows_total", {"outcome": "rejected"})

    upload(upload_client, workbook_bytes([row("INV-1"), row("INV-2", amount=0)]))

    assert sample("invoice_uploads_total", {"result": "success"}) == before_uploads + 1
    assert sample("invoice_rows_total", {"outcome": "rejected"}) == before_rejected + 1
    completed = next(line for line in logs() if line["event"] == "invoice_upload.completed")
    assert (completed["rows_created"], completed["rows_rejected"]) == (1, 1)
    assert "request_id" in completed
    assert "INV-1" not in str(completed)  # row contents are never logged


# --- AC10: performance ---------------------------------------------------------------------


@pytest.mark.slow
def test_INV_001_AC10_5000_rows_in_under_10_seconds(upload_client):
    start_date = date(2026, 1, 1)
    rows = [
        row(
            f"PERF-{i}",
            f"Customer {i % 97}",
            start_date,
            start_date + timedelta(days=30 + i % 60),
            1000 + i,
            start_date + timedelta(days=10) if i % 3 == 0 else None,
        )
        for i in range(5000)
    ]
    content = workbook_bytes(rows)

    started = time.perf_counter()
    body = upload(upload_client, content).json()
    elapsed = time.perf_counter() - started

    assert body["rows_created"] == 5000
    assert elapsed < 10, f"took {elapsed:.1f}s"
