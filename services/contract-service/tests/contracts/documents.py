"""Test documents built in code (no binary fixtures in the repo)."""

from io import BytesIO

import docx
from fpdf import FPDF

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF = "application/pdf"

FILLER = (
    "The Supplier shall deliver the goods described in Annexure 1 within fourteen days of each "
    "purchase order. The Buyer shall pay each invoice within thirty days of the invoice date. "
)


def docx_bytes(*paragraphs: str, table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def contract_docx(title: str = "Supply Agreement", unique: str = "") -> bytes:
    """A plausible contract with enough text to read; `unique` makes the file hash differ."""
    return docx_bytes(
        title,
        "This Supply Agreement is made on 1 April 2026 between Sharma Textiles Private Limited "
        '(the "Supplier") and Deccan Printing Works (the "Buyer").',
        "This Agreement is effective from 1 April 2026 and shall continue until 31 March 2027.",
        FILLER * 2,
        "The Client shall indemnify the Service Provider against all losses without limit.",
        f"Reference {unique}" if unique else "Reference none",
    )


def text_pdf_bytes(*lines: str) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in lines:
        pdf.multi_cell(0, 6, line, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def image_only_pdf_bytes() -> bytes:
    """Like a scanned page: drawing only, no text layer."""
    pdf = FPDF()
    pdf.add_page()
    pdf.set_fill_color(40, 40, 40)
    for row in range(20):
        pdf.rect(20, 20 + row * 10, 170, 4, style="F")
    return bytes(pdf.output())
